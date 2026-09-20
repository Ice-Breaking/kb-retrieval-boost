# -*- coding: utf-8 -*-
"""把原始政策文件清洗、切分成结构化知识切片。

输入: data/raw/ 下的 .docx / .doc / .pdf / .txt（该目录不入库）
输出: data/knowledge_base.json

用法:
    python src/build_corpus.py
    python src/build_corpus.py --dump-index /tmp/chunk_index.txt
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
OUT = ROOT / "data" / "knowledge_base.json"

# data/raw 下的文件清单（file 为不含扩展名的文件名）
DOCS = [
    dict(doc_id="tj_subsidy_2022", file="天津市职业技能培训补贴实施办法",
         title="天津市职业技能培训补贴实施办法", doc_no="津人社局发〔2022〕20号",
         doc_type="市级办法", year=2022, splitter="article"),
    dict(doc_id="tj_supervise_2022", file="人社办发（2022）60号(1)",
         title="天津市职业技能培训监督管理暂行办法", doc_no="津人社办发〔2022〕60号",
         doc_type="市级办法", year=2022, splitter="article"),
    dict(doc_id="tj_ops_manual", file="职业培训补贴经办流程（最终稿）",
         title="天津市职业培训补贴经办操作规程", doc_no="内部参考",
         doc_type="经办规程", year=2019, splitter="section"),
    dict(doc_id="tj_catalog_2024", file="市人社局市财政局关于印发《天津市市场紧缺职业需求程度及培训补贴标准目录（2024年版）》的通知",
         title="天津市市场紧缺职业需求程度及培训补贴标准目录（2024年版）", doc_no="津人社办发〔2023〕60号",
         doc_type="补贴目录", year=2024, splitter="notice", version="2024年版"),
    dict(doc_id="tj_catalog_2026", file="市人社局市财政局关于印发《天津市市场紧缺职业需求程度及培训补贴标准目录（2026年版）》的通知",
         title="天津市市场紧缺职业需求程度及培训补贴标准目录（2026年版）", doc_no="津人社办发〔2026〕10号",
         doc_type="补贴目录", year=2026, splitter="notice", version="2026年版"),
    dict(doc_id="national_fund_2023", file="就业补助资金管理办法",
         title="就业补助资金管理办法", doc_no="财社〔2023〕181号",
         doc_type="国家级办法", year=2023, splitter="article", pdf_noise=True),
]

ART_RE = re.compile(r"第([一二三四五六七八九十百]+)条")
CHAP_RE = re.compile(r"第[一二三四五六七八九十]+章")


# ---------------------------------------------------------------- 读取
def read_text(path: Path) -> str:
    """按扩展名读取 docx / doc / pdf / txt 的纯文本。"""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        import docx
        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            parts.append("[表格]")
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                dedup = [cells[0]] if cells else []
                for cell in cells[1:]:
                    if cell and cell != dedup[-1]:
                        dedup.append(cell)
                line = " | ".join(dedup).strip(" |")
                if line:
                    parts.append(line)
        return "\n".join(parts)
    if suffix == ".doc":
        proc = subprocess.run(
            ["textutil", "-convert", "txt", "-encoding", "UTF-8", "-stdout", str(path)],
            capture_output=True, text=True)
        if proc.returncode != 0 or not proc.stdout.strip():
            sys.exit(f"读取 .doc 失败（macOS 依赖 textutil）：{path}\n{proc.stderr[:200]}")
        return proc.stdout
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"不支持的文件类型：{path}")


# ---------------------------------------------------------------- 清洗
def clean_text(text: str, pdf_noise: bool = False) -> str:
    """通用清洗：去页码行、压缩空白、去掉乱码。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # 去掉独占一行的页码（形如 "1" / "-2" / "— 3 —"）
    text = re.sub(r"(?m)^\s*[-—–]?\s*\d{1,3}\s*[-—–]?\s*$", "", text)
    # 去掉 Word 域代码残留
    text = re.sub(r"—\s*PAGE\s*\\?\*?\s*MERGEFORMAT\s*\d*\s*—", "", text)
    if pdf_noise:
        # 双栏 PDF 抽取结果是碎片化的：去掉全部空白后再按句读补回切分点
        text = re.sub(r"\s+", "", text)
        text = text.replace("\x00", "").replace("９", "，")
        text = text.replace("，", "，\n").replace("。", "。\n")
    lines = []
    for line in text.split("\n"):
        line = re.sub(r"[ \t]+", " ", line).replace("\u3000", " ").strip()
        if not line:
            continue
        if re.fullmatch(r"[-—–\s]*\d{1,3}[-—–\s]*", line):  # 纯页码行
            continue
        lines.append(line)
    return "\n".join(lines)


def split_long(body: str, max_len: int) -> list:
    """超长文本按括号条目、再按句读切分。"""
    if len(body) <= max_len:
        return [body]
    parts, buf = [], ""
    for piece in re.split(r"(?=（[一二三四五六七八九十]+）)", body):
        if len(buf) + len(piece) <= max_len:
            buf += piece
        else:
            if buf:
                parts.append(buf)
            while len(piece) > max_len:
                cut = max(piece.rfind("；", 0, max_len), piece.rfind("。", 0, max_len))
                cut = cut + 1 if cut > max_len // 2 else max_len
                parts.append(piece[:cut])
                piece = piece[cut:]
            buf = piece
    if buf:
        parts.append(buf)
    return parts


# ---------------------------------------------------------------- 切分
def chunk_articles(text: str, max_len: int) -> list:
    """条文型文档：按“第X条”切分，并记录所属章节与前置通知。"""
    chapters = [(m.start(), re.sub(r"\s+", "", m.group(0))) for m in
                re.finditer(r"第[一二三四五六七八九十]+章[^\n]{0,12}", text)]
    marks = [(m.start(), m.group(1)) for m in ART_RE.finditer(text)]
    out = []
    if marks and marks[0][0] > 30:
        for i, piece in enumerate(split_long(text[:marks[0][0]].strip(), max_len)):
            out.append(dict(locator="前言" if i == 0 else f"前言（续{i + 1}）",
                            chapter="", part="通知", content=piece))
    for idx, (pos, num) in enumerate(marks):
        end = marks[idx + 1][0] if idx + 1 < len(marks) else len(text)
        body = text[pos:end].strip()
        chapter = ""
        for cpos, cname in chapters:
            if cpos < pos:
                chapter = cname
        for i, piece in enumerate(split_long(body, max_len)):
            out.append(dict(locator=f"第{num}条" + ("" if i == 0 else f"（续{i + 1}）"),
                            chapter=chapter, part="正文", content=piece))
    return out


def chunk_sections(text: str, max_len: int) -> list:
    """规程/流程型文档：按“一、二、三、”大标题切分，超长再细分。"""
    out, title, buf = [], "前言", []

    def flush():
        nonlocal buf
        body = "\n".join(line for line in buf if line.strip()).strip()
        if body:
            for i, piece in enumerate(split_long(body, max_len)):
                loc = title if i == 0 else f"{title}（续{i + 1}）"
                out.append(dict(locator=loc, chapter="", part="正文", content=piece))
        buf = []

    for line in text.split("\n"):
        if re.fullmatch(r"[一二三四五六七八九十]+、.{0,30}", line):
            flush()
            title = line
        else:
            buf.append(line)
    flush()
    return out


def chunk_notice(text: str, max_len: int) -> list:
    """通知/目录型文档：正文按段落成块，附表单独成块。"""
    head, _, table = text.partition("[表格]")
    out, buf = [], ""
    for line in head.split("\n"):
        if len(buf) + len(line) + 1 <= max_len:
            buf += line + "\n"
        else:
            if buf.strip():
                out.append(buf.strip())
            buf = line + "\n"
    if buf.strip():
        out.append(buf.strip())
    for i, piece in enumerate(out):
        out[i] = dict(locator="通知正文" if i == 0 else f"通知正文（续{i + 1}）",
                      chapter="", part="通知", content=piece)
    if table.strip():
        lines = [re.sub(r"\s*\|\s*", " | ", line).strip(" |") for line in table.split("\n")]
        dedup = [lines[0]] if lines else []
        for line in lines[1:]:
            if line and line != dedup[-1]:
                dedup.append(line)
        body = "\n".join(dedup)
        for i, piece in enumerate(split_long(body, 1000)):
            out.append(dict(locator="补贴标准表" + ("" if i == 0 else f"（续{i + 1}）"),
                            chapter="", part="附表", content=piece))
    return out


# ---------------------------------------------------------------- 主流程
def build(raw_dir: Path, out_path: Path, dump_index: str = None) -> list:
    records = []
    for doc in DOCS:
        matches = sorted(raw_dir.glob(doc["file"] + ".*"))
        if not matches:
            sys.exit(f"缺少原始文件：{raw_dir}/{doc['file']}.*")
        text = clean_text(read_text(matches[0]), pdf_noise=doc.get("pdf_noise", False))
        max_len = 900 if doc["splitter"] == "article" else (700 if doc["splitter"] == "section" else 450)
        if doc["splitter"] == "article":
            items = chunk_articles(text, max_len)
        elif doc["splitter"] == "section":
            items = chunk_sections(text, max_len)
        else:
            items = chunk_notice(text, max_len)
        kept = [it for it in items if len(it["content"].strip()) >= 30]
        for i, item in enumerate(kept, 1):
            content = re.sub(r"\n{2,}", "\n", item["content"]).strip()
            records.append({
                "id": f"{doc['doc_id']}#{i:03d}",
                "doc_id": doc["doc_id"],
                "doc_title": doc["title"],
                "doc_no": doc["doc_no"],
                "doc_type": doc["doc_type"],
                "year": doc["year"],
                "version": doc.get("version", ""),
                "chapter": item.get("chapter", ""),
                "locator": item["locator"],
                "part": item.get("part", "正文"),
                "source_label": f"《{doc['title']}》{item['locator']}",
                "content": content,
                "n_chars": len(content),
            })
        print(f"  {doc['doc_id']:<22} 切片 {len(kept):>3} 个  原文 {len(text):>6} 字")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(records, ensure_ascii=False, indent=1), encoding="utf-8")
    total = sum(r["n_chars"] for r in records)
    print(f"\n共 {len(records)} 个切片，合计 {total} 字 -> {out_path}")
    if dump_index:
        lines = [f"{r['id']}  [{r['n_chars']:>4}字]  {r['source_label']}\n      {r['content'][:100]}"
                 for r in records]
        Path(dump_index).write_text("\n".join(lines), encoding="utf-8")
        print(f"切片索引已导出 -> {dump_index}")
    return records


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="构建知识切片")
    parser.add_argument("--raw", default=str(RAW))
    parser.add_argument("--out", default=str(OUT))
    parser.add_argument("--dump-index", default=None)
    args = parser.parse_args()
    print("开始构建知识切片：")
    build(Path(args.raw), Path(args.out), args.dump_index)
